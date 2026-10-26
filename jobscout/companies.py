"""Load the target company list from the Excel workbook into the DB.

The single input you maintain. Two load modes:

* ``load_from_excel``  — every company on the named sheets.
* ``load_yellow_from_excel`` — only the rows the user highlighted **yellow**
  in column B (fill ``FFFFFF00``). This is the canonical "target shortlist":
  the companies the recommender is expected to cover end to end.
"""
from __future__ import annotations

import csv
from pathlib import Path
from typing import Iterable, Optional

import openpyxl

from . import config, db

POC_SHEETS = ("Austrian Tier-1 & 2", "Big Tech (US-HQ EU offices)")
_HEADER_MARKER = "Organization / Company"
# Pure-yellow highlight the user paints on their target rows.
YELLOW_FILL = "FFFFFF00"
NAME_COL = 2          # column B — "Organization / Company"
LOC_COL = 3           # column C — "Locations to look out for"
PCT_COL = 4           # column D — "% Interview"


def load_from_excel(xlsx_path: Path, sheets: Iterable[str] = POC_SHEETS) -> list[dict]:
    wb = openpyxl.load_workbook(str(xlsx_path), read_only=True, data_only=True)
    rows: list[dict] = []
    for sheet in sheets:
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        header_row = None
        for r in ws.iter_rows(values_only=True):
            cells = [str(c).strip() if c is not None else "" for c in r]
            if header_row is None:
                if any(c == _HEADER_MARKER for c in cells):
                    header_row = True
                continue
            # data rows: [#, name, locations, % interview, last visited]
            if len(cells) < 2 or not cells[1]:
                continue
            name = cells[1]
            locations = cells[2] if len(cells) > 2 else ""
            pct = None
            if len(cells) > 3 and cells[3]:
                try:
                    pct = float(cells[3])
                except ValueError:
                    pct = None
            rows.append({
                "name": name,
                "category": sheet,
                "locations": locations,
                "interview_pct": pct,
            })
    wb.close()
    return rows


def _is_yellow(cell) -> bool:
    f = cell.fill
    return bool(f and f.patternType == "solid"
                and str(getattr(f.fgColor, "rgb", None)) == YELLOW_FILL)


def load_yellow_from_excel(xlsx_path: Path,
                           sheets: Optional[Iterable[str]] = None) -> list[dict]:
    """Every row whose company cell (col B) is highlighted yellow.

    Reads cell *fills*, so the workbook is opened with styles (not read_only).
    """
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=True)
    sheet_names = list(sheets) if sheets else wb.sheetnames
    rows: list[dict] = []
    for sheet in sheet_names:
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        for r in range(1, ws.max_row + 1):
            name_cell = ws.cell(r, NAME_COL)
            if not _is_yellow(name_cell):
                continue
            name = (str(name_cell.value).strip() if name_cell.value else "")
            if not name:
                continue
            loc = ws.cell(r, LOC_COL).value
            pct_val = ws.cell(r, PCT_COL).value
            try:
                pct = float(pct_val) if pct_val not in (None, "") else None
            except (ValueError, TypeError):
                pct = None
            rows.append({
                "name": name,
                "category": sheet,
                "locations": str(loc).strip() if loc else "",
                "interview_pct": pct,
            })
    wb.close()
    return rows


def _seed_rows(rows: list[dict]) -> int:
    for row in rows:
        db.upsert_company(
            name=row["name"],
            category=row["category"],
            locations=row["locations"],
            interview_pct=row["interview_pct"],
        )
    # Also drop a CSV snapshot so the list is easy to view/edit outside the DB.
    with open(config.COMPANIES_SEED_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["name", "category", "locations", "interview_pct"])
        w.writeheader()
        w.writerows(rows)
    return len(rows)


def seed_companies(xlsx_path: Path, sheets: Iterable[str] = POC_SHEETS) -> int:
    return _seed_rows(load_from_excel(xlsx_path, sheets))


def seed_yellow_companies(xlsx_path: Path) -> int:
    """Seed exactly the user's yellow-highlighted target shortlist."""
    return _seed_rows(load_yellow_from_excel(xlsx_path))
